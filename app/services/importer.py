"""Importador único de la hoja de cálculo original.

La app es la fuente de verdad: esto se usa una vez, para traer el histórico
que vivía en la hoja de Google Sheets (o en cualquier export suyo). Acepta
`.xlsx` y `.csv`, sin credenciales ni acceso a la red.

El histórico vive repartido en dos sitios dentro del mismo libro:

- **La tabla de datos** (pestaña `Kikis`) trae los encuentros con su pretexto,
  motivación, valoraciones y observaciones. No tiene columna de tipo: usa
  `calidad` y `tiempo` a 0 para marcar un desencuentro, y nunca deja sólo una
  de las dos a cero, así que la regla 0/0 → "No Kiki" se deduce del propio
  fichero. Si algún día la hoja trae columna de tipo, esa manda.
- **Las pestañas de calendario** (una por año) marcan la menstruación pintando
  la celda del día de rojo. Esos días no están en la tabla, así que se leen
  del color de relleno.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta
from typing import Any, Iterable, Optional

from app import crud, db
from app.schemas import MOTIVACIONES, TIPOS_CON_CONTEXTO, TIPOS_CON_TAGS
from app.schemas import TIPOS_PUNTUADOS
from app.schemas import TIPO_GAYOLA, TIPO_KIKI, TIPO_MAREA, TIPO_NO_KIKI

MAX_BYTES = 10 * 1024 * 1024
META_ULTIMA_IMPORTACION = "import.ultima"

# Alias de cabecera aceptados, comparados sin acentos ni mayúsculas.
ALIAS_COLUMNA: dict[str, str] = {
    "fecha": "fecha", "dia": "fecha", "date": "fecha",
    "tipo": "tipo", "categoria": "tipo", "evento": "tipo",
    "pretexto": "pretexto", "motivo": "pretexto", "excusa": "pretexto",
    "motivacion": "motivacion", "iniciativa": "motivacion", "quien": "motivacion",
    "calidad": "calidad", "nota": "calidad", "valoracion": "calidad",
    "tiempo": "tiempo", "duracion": "tiempo",
    "observaciones": "observaciones", "comentarios": "observaciones",
    "notas": "observaciones", "observacion": "observaciones",
    "tags": "tags", "etiquetas": "tags", "practicas": "tags", "posturas": "tags",
}

ALIAS_TIPO: dict[str, str] = {
    "kiki": TIPO_KIKI, "si": TIPO_KIKI, "encuentro": TIPO_KIKI,
    "nokiki": TIPO_NO_KIKI, "no": TIPO_NO_KIKI, "desencuentro": TIPO_NO_KIKI,
    "gayola": TIPO_GAYOLA,
    "marea": TIPO_MAREA, "regla": TIPO_MAREA, "periodo": TIPO_MAREA,
    "menstruacion": TIPO_MAREA,
}

# Origen del contador de fechas de las hojas de cálculo.
EPOCA_HOJA = date(1899, 12, 30)

# Una cabecera válida necesita la fecha y al menos otra columna reconocible:
# así se descartan las pestañas de calendario, llenas de celdas sueltas.
MINIMO_COLUMNAS_CABECERA = 2


class ImportError_(ValueError):
    """El fichero no se puede leer o no contiene datos reconocibles."""


# --- Normalización ---------------------------------------------------------

def _slug(valor: Any) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", texto.lower())


def normalize_tipo(valor: Any) -> Optional[str]:
    return ALIAS_TIPO.get(_slug(valor))


def normalize_motivacion(valor: Any) -> Optional[str]:
    objetivo = _slug(valor)
    for motivacion in MOTIVACIONES:
        if _slug(motivacion) == objetivo:
            return motivacion
    return None


def normalize_fecha(valor: Any) -> Optional[str]:
    """Acepta datetime, ISO, dd/mm/aaaa y el número de serie de la hoja."""
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()

    texto = str(valor).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", texto):
        return texto
    if re.fullmatch(r"\d{1,6}(\.0+)?", texto):
        return (EPOCA_HOJA + timedelta(days=int(float(texto)))).isoformat()

    coincidencia = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", texto)
    if coincidencia:
        dia, mes, anio = (int(parte) for parte in coincidencia.groups())
        if anio < 100:
            anio += 2000
        try:
            return date(anio, mes, dia).isoformat()
        except ValueError:
            return None
    return None


def _puntuacion(valor: Any) -> int:
    try:
        numero = int(float(str(valor).strip().replace(",", ".")))
    except (TypeError, ValueError):
        return 0
    return max(0, min(4, numero))


def _texto(valor: Any) -> Optional[str]:
    limpio = str(valor).strip() if valor is not None else ""
    # "N/A" en la hoja original significa "sin pretexto".
    if not limpio or limpio.upper() in {"N/A", "NA", "-", "—"}:
        return None
    return limpio


def _tags(valor: Any) -> list[str]:
    """Una celda de etiquetas viene separada por comas, punto y coma o barras."""
    texto = _texto(valor)
    if not texto:
        return []
    partes = re.split(r"[,;/|]", texto)
    return [parte.strip() for parte in partes if parte.strip()]


def inferir_tipo(calidad: int, tiempo: int, explicito: Optional[str]) -> str:
    """Un 0 en ambas valoraciones marca el desencuentro en la hoja original."""
    if explicito:
        return explicito
    return TIPO_NO_KIKI if calidad == 0 and tiempo == 0 else TIPO_KIKI


# --- Lectura ---------------------------------------------------------------

def _mapear_cabecera(fila: Iterable[Any]) -> dict[int, str]:
    """Devuelve {índice de columna: campo} para las columnas reconocidas."""
    mapa: dict[int, str] = {}
    for indice, celda in enumerate(fila):
        campo = ALIAS_COLUMNA.get(_slug(celda))
        if campo and campo not in mapa.values():
            mapa[indice] = campo
    return mapa


def _es_cabecera(mapa: dict[int, str]) -> bool:
    return "fecha" in mapa.values() and len(mapa) >= MINIMO_COLUMNAS_CABECERA


def fila_a_registro(valores: list[Any], mapa: dict[int, str]) -> Optional[dict[str, Any]]:
    """Convierte una fila cruda al esquema interno, o None si no es un registro."""
    crudo = {
        campo: valores[indice]
        for indice, campo in mapa.items()
        if indice < len(valores)
    }

    fecha = normalize_fecha(crudo.get("fecha"))
    if not fecha:
        return None

    calidad = _puntuacion(crudo.get("calidad"))
    tiempo = _puntuacion(crudo.get("tiempo"))
    tipo = inferir_tipo(calidad, tiempo, normalize_tipo(crudo.get("tipo")))
    # Las mismas reglas que aplica el modelo: sólo el Kiki se valora, y sólo
    # Kiki y No Kiki llevan pretexto y motivación.
    puntuado = tipo in TIPOS_PUNTUADOS
    con_contexto = tipo in TIPOS_CON_CONTEXTO
    con_tags = tipo in TIPOS_CON_TAGS

    return {
        "fecha": fecha,
        "tipo": tipo,
        "pretexto": _texto(crudo.get("pretexto")) if con_contexto else None,
        "motivacion": normalize_motivacion(crudo.get("motivacion")) if con_contexto else None,
        "calidad": calidad if puntuado else 0,
        "tiempo": tiempo if puntuado else 0,
        "observaciones": _texto(crudo.get("observaciones")),
        "tags": _tags(crudo.get("tags")) if con_tags else [],
    }


def _leer_tabla(filas: list[list[Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Localiza la cabecera y convierte el resto. Devuelve (registros, descartes)."""
    mapa: dict[int, str] = {}
    inicio = 0
    # La cabecera no siempre está en la primera fila: puede haber títulos o
    # celdas sueltas encima.
    for numero, fila in enumerate(filas[:20]):
        candidato = _mapear_cabecera(fila)
        if _es_cabecera(candidato):
            mapa, inicio = candidato, numero + 1
            break

    if not mapa:
        return [], []

    registros: list[dict[str, Any]] = []
    descartes: list[dict[str, Any]] = []
    for numero, fila in enumerate(filas[inicio:], start=inicio + 1):
        if not any(celda not in (None, "") for celda in fila):
            continue
        registro = fila_a_registro(list(fila), mapa)
        if registro:
            registros.append(registro)
        else:
            resumen = " | ".join(str(c) for c in fila if c not in (None, ""))[:120]
            descartes.append({"fila": numero, "motivo": "sin fecha válida",
                              "contenido": resumen})
    return registros, descartes


def _abrir_libro(contenido: bytes) -> Any:
    """Abre el .xlsx con estilos: el color de relleno es un dato, no adorno."""
    try:
        import openpyxl
    except ImportError as error:  # pragma: no cover - depende del entorno
        raise ImportError_("Falta la librería openpyxl para leer ficheros .xlsx") from error

    try:
        return openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
    except Exception as error:
        raise ImportError_(f"No se pudo abrir el fichero Excel: {error}") from error


def _hojas_csv(contenido: bytes) -> list[tuple[str, list[list[Any]]]]:
    texto = contenido.decode("utf-8-sig", errors="replace")
    muestra = texto[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    filas = [list(fila) for fila in csv.reader(io.StringIO(texto), dialecto)]
    return [("csv", filas)]


def analizar(contenido: bytes, nombre: str,
             hoja: Optional[str] = None) -> dict[str, Any]:
    """Lee el fichero y devuelve los registros encontrados, sin tocar la BBDD.

    Junta dos fuentes del mismo libro: la tabla de datos —de las varias
    pestañas se queda con la que más registros válidos produce, así los
    calendarios se descartan solos— y los días de Marea pintados de rojo en
    las pestañas de calendario.
    """
    if not contenido:
        raise ImportError_("El fichero está vacío")
    if len(contenido) > MAX_BYTES:
        raise ImportError_(f"El fichero supera el límite de {MAX_BYTES // 1024 // 1024} MB")

    minuscula = nombre.lower()
    if minuscula.endswith((".xlsx", ".xlsm")):
        libro = _abrir_libro(contenido)
        hojas = [
            (titulo, [list(fila) for fila in libro[titulo].iter_rows(values_only=True)])
            for titulo in libro.sheetnames
        ]
        mareas, colores_ignorados = leer_mareas(libro)
    elif minuscula.endswith((".csv", ".tsv", ".txt")):
        hojas = _hojas_csv(contenido)
        mareas, colores_ignorados = [], {}
    else:
        raise ImportError_(f"Formato no soportado: {nombre}. Usa .xlsx o .csv")

    if hoja:
        hojas = [par for par in hojas if par[0] == hoja]
        if not hojas:
            raise ImportError_(f"El fichero no tiene ninguna pestaña llamada «{hoja}»")

    mejor_nombre, mejor_registros, mejor_descartes = None, [], []
    for nombre_hoja, filas in hojas:
        registros, descartes = _leer_tabla(filas)
        if len(registros) > len(mejor_registros):
            mejor_nombre, mejor_registros, mejor_descartes = nombre_hoja, registros, descartes

    todos = mejor_registros + mareas
    if not todos:
        raise ImportError_(
            "No se ha encontrado ninguna tabla con columna de fecha ni días "
            "marcados en los calendarios. Revisa que la hoja tenga cabecera."
        )

    fechas = sorted(registro["fecha"] for registro in todos)
    por_tipo = {tipo: 0 for tipo in (TIPO_KIKI, TIPO_NO_KIKI, TIPO_MAREA)}
    for registro in todos:
        por_tipo[registro["tipo"]] += 1

    return {
        "origen": nombre,
        "hoja_datos": mejor_nombre,
        "hojas_disponibles": [par[0] for par in hojas],
        "registros": todos,
        "mareas_del_calendario": len(mareas),
        "descartadas": mejor_descartes[:20],
        "total_descartadas": len(mejor_descartes),
        "colores_ignorados": colores_ignorados,
        "por_tipo": por_tipo,
        "rango": {"desde": fechas[0], "hasta": fechas[-1]},
    }


# --- Calendarios por color -------------------------------------------------

MESES_CALENDARIO = (
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
)

# Alto máximo de un bloque de mes bajo su título, en filas.
ALTO_BLOQUE = 8
DIAS_SEMANA = 7

# Umbrales para clasificar el relleno de una celda.
GRIS_MAX_CROMA = 0.06      # por debajo es blanco o gris: celda vacía
ROJO_ARCO = 20.0           # grados de tono a cada lado del rojo puro


def _tono_y_croma(rgb: str) -> Optional[tuple[float, float]]:
    """Tono en grados y croma (0-1) de un color ARGB/RGB hexadecimal."""
    limpio = (rgb or "").strip()
    if len(limpio) == 8:      # ARGB
        limpio = limpio[2:]
    if len(limpio) != 6:
        return None
    try:
        rojo, verde, azul = (int(limpio[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except ValueError:
        return None

    alto, bajo = max(rojo, verde, azul), min(rojo, verde, azul)
    croma = alto - bajo
    if croma == 0:
        return 0.0, 0.0
    if alto == rojo:
        tono = 60 * (((verde - azul) / croma) % 6)
    elif alto == verde:
        tono = 60 * (((azul - rojo) / croma) + 2)
    else:
        tono = 60 * (((rojo - verde) / croma) + 4)
    return tono, croma


def es_rojo(rgb: Optional[str]) -> bool:
    """¿El relleno es de la familia del rojo? Incluye los tonos claros.

    Se clasifica por tono y no por hex exacto para que valga igual un rojo
    pleno que un rosa de manchado.
    """
    if not rgb:
        return False
    medida = _tono_y_croma(rgb)
    if medida is None:
        return False
    tono, croma = medida
    if croma < GRIS_MAX_CROMA:
        return False
    return tono <= ROJO_ARCO or tono >= 360 - ROJO_ARCO


def _anio_de_hoja(nombre: str, filas: list[list[Any]]) -> Optional[int]:
    if re.fullmatch(r"(19|20)\d{2}", nombre.strip()):
        return int(nombre.strip())
    for fila in filas[:6]:
        for celda in fila:
            texto = str(celda or "").strip()
            if re.fullmatch(r"(19|20)\d{2}", texto):
                return int(texto)
    return None


def _celdas_de_calendario(hoja: Any, anio: int) -> Iterable[tuple[date, Any]]:
    """Recorre los bloques de mes del calendario y devuelve (fecha, celda)."""
    slugs_mes = [_slug(nombre) for nombre in MESES_CALENDARIO]

    bloques: set[tuple[int, int, int]] = set()
    for fila in hoja.iter_rows():
        for celda in fila:
            etiqueta = _slug(celda.value)
            if etiqueta in slugs_mes:
                bloques.add((slugs_mes.index(etiqueta) + 1, celda.row, celda.column))

    for mes, fila_mes, columna_mes in sorted(bloques):
        for desplazamiento in range(1, ALTO_BLOQUE + 1):
            for columna in range(DIAS_SEMANA):
                celda = hoja.cell(row=fila_mes + desplazamiento,
                                  column=columna_mes + columna)
                if not isinstance(celda.value, (int, float)):
                    continue
                dia = int(celda.value)
                if not 1 <= dia <= 31:
                    continue
                try:
                    yield date(anio, mes, dia), celda
                except ValueError:
                    continue


def _relleno(celda: Any) -> Optional[str]:
    relleno = celda.fill
    if relleno is None or relleno.patternType != "solid":
        return None
    color = relleno.fgColor
    return color.rgb if getattr(color, "type", None) == "rgb" else None


def leer_mareas(libro: Any) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Extrae los días de Marea de las pestañas de calendario.

    Devuelve también el recuento de colores no reconocidos, para que el
    informe pueda decir qué se ha dejado fuera en lugar de callarlo.
    """
    mareas: dict[str, dict[str, Any]] = {}
    ignorados: dict[str, int] = {}

    for nombre in libro.sheetnames:
        hoja = libro[nombre]
        filas = [list(fila) for fila in hoja.iter_rows(max_row=6, values_only=True)]
        anio = _anio_de_hoja(nombre, filas)
        if anio is None:
            continue
        for fecha, celda in _celdas_de_calendario(hoja, anio):
            color = _relleno(celda)
            if es_rojo(color):
                mareas[fecha.isoformat()] = {
                    "fecha": fecha.isoformat(),
                    "tipo": TIPO_MAREA,
                    "pretexto": None,
                    "motivacion": None,
                    "calidad": 0,
                    "tiempo": 0,
                    "observaciones": None,
                }
            elif color:
                medida = _tono_y_croma(color)
                if medida and medida[1] >= GRIS_MAX_CROMA:
                    ignorados[color] = ignorados.get(color, 0) + 1

    return sorted(mareas.values(), key=lambda registro: registro["fecha"]), ignorados


# --- Escritura -------------------------------------------------------------

# Campos que identifican un registro. No incluye el id: sirve para emparejar
# lo que llega del fichero con lo que ya está guardado.
CAMPOS_IDENTIDAD = (
    "fecha", "tipo", "pretexto", "motivacion", "calidad", "tiempo", "observaciones",
)


def _huella(registro: dict[str, Any]) -> tuple[Any, ...]:
    """Las etiquetas no entran: se pueden editar sin dejar de ser el mismo evento."""
    return tuple(registro.get(campo) for campo in CAMPOS_IDENTIDAD)


def importar(contenido: bytes, nombre: str, modo: str = "combinar",
             simular: bool = False, hoja: Optional[str] = None) -> dict[str, Any]:
    """Analiza el fichero y, si no es simulacro, vuelca los registros.

    - `combinar`: inserta lo que falta y deja intacto lo que ya existe.
      Repetirlo no duplica nada.
    - `reemplazar`: vacía la tabla antes de insertar.
    """
    analisis = analizar(contenido, nombre, hoja)
    registros = analisis.pop("registros")

    # Un día admite varios eventos, incluso del mismo tipo (dos Kikis en la
    # misma fecha). Por eso el emparejamiento es por registro completo y
    # cuenta repeticiones: si la BBDD tiene uno y el fichero trae dos, entra
    # el que falta en lugar de descartarse ambos.
    disponibles: Counter[tuple[Any, ...]] = Counter()
    if modo == "combinar":
        actuales, _ = crud.list_events()
        disponibles = Counter(_huella(evento) for evento in actuales)

    nuevos = []
    for registro in registros:
        huella = _huella(registro)
        if disponibles[huella]:
            disponibles[huella] -= 1
        else:
            nuevos.append(registro)

    informe = analisis | {
        "modo": modo,
        "simulado": simular,
        "leidas": len(registros),
        "nuevas": len(nuevos),
        "ya_existentes": len(registros) - len(nuevos),
        "muestra": registros[:5],
    }

    if simular:
        informe["insertadas"] = 0
        return informe

    if modo == "reemplazar":
        with db.connect() as conexion:
            borradas = conexion.execute("DELETE FROM events").rowcount
        informe["borradas"] = borradas
        nuevos = registros

    insertadas, _ = crud.upsert_many([dict(registro, id=None) for registro in nuevos])
    informe["insertadas"] = insertadas
    db.set_meta(
        META_ULTIMA_IMPORTACION,
        f"{datetime.now().isoformat(timespec='seconds')} · {nombre} · "
        f"{insertadas} registros",
    )
    return informe


def ultima_importacion() -> Optional[str]:
    return db.get_meta(META_ULTIMA_IMPORTACION)
