"""Importador: normalización, deducción de tipo y lectura de calendarios."""

from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from app.services import importer
from app.services.importer import ImportError_


# --- Normalización ---------------------------------------------------------

@pytest.mark.parametrize(
    "entrada, esperado",
    [
        ("2026-03-10", "2026-03-10"),
        ("10/03/2026", "2026-03-10"),
        ("10-3-26", "2026-03-10"),
        ("46091", "2026-03-10"),          # número de serie de la hoja
        ("46091.0", "2026-03-10"),
        (date(2026, 3, 10), "2026-03-10"),
        ("", None),
        ("no es fecha", None),
        ("32/13/2026", None),
    ],
)
def test_normalize_fecha(entrada, esperado):
    assert importer.normalize_fecha(entrada) == esperado


@pytest.mark.parametrize(
    "entrada, esperado",
    [("Propia", "Propia"), ("propia ", "Propia"), ("AMBOS", "Ambos"),
     ("ajena", "Ajena"), ("nadie", None)],
)
def test_normalize_motivacion(entrada, esperado):
    assert importer.normalize_motivacion(entrada) == esperado


def test_na_cuenta_como_sin_pretexto():
    assert importer._texto("N/A") is None
    assert importer._texto("  ") is None
    assert importer._texto(" Calentura ") == "Calentura"


def test_tipo_se_deduce_de_las_valoraciones():
    # La hoja original marca el desencuentro poniendo ambas a cero.
    assert importer.inferir_tipo(0, 0, None) == "No Kiki"
    assert importer.inferir_tipo(3, 2, None) == "Kiki"
    assert importer.inferir_tipo(0, 1, None) == "Kiki"
    # Una columna de tipo explícita manda sobre la deducción.
    assert importer.inferir_tipo(0, 0, "Marea") == "Marea"


# --- Clasificación de color ------------------------------------------------

@pytest.mark.parametrize(
    "color, rojo",
    [
        ("FFFF0000", True),    # rojo pleno
        ("FFCC0000", True),    # rojo oscuro
        ("FFE06666", True),    # rojo medio
        ("FFF4CCCC", True),    # rosa claro
        ("FF00FF00", False),   # verde
        ("FFFF00FF", False),   # magenta
        ("FFFFFFCC", False),   # amarillo claro
        ("FFFFFFFF", False),   # blanco
        ("FFDDDDDD", False),   # gris
        (None, False),
    ],
)
def test_es_rojo(color, rojo):
    assert importer.es_rojo(color) is rojo


# --- Lectura de ficheros ---------------------------------------------------

def _libro_de_prueba() -> bytes:
    libro = openpyxl.Workbook()

    datos = libro.active
    datos.title = "Kikis"
    datos.append(["Fecha", "Pretexto", "Motivación", "Calidad", "Tiempo", "Observaciones"])
    datos.append([date(2026, 1, 5), "Calentura", "Propia ", 3, 2, " buena "])
    datos.append([date(2026, 1, 9), "N/A", "Ajena", 0, 0, None])

    calendario = libro.create_sheet("2026")
    calendario.cell(row=3, column=3, value="ENERO")
    for columna, dia in enumerate(["L", "M", "M", "J", "V", "S", "D"]):
        calendario.cell(row=4, column=3 + columna, value=dia)
    for columna in range(7):
        celda = calendario.cell(row=5, column=3 + columna, value=12 + columna)
        if columna < 3:
            celda.fill = PatternFill("solid", fgColor="FFFF0000")   # marea
        elif columna == 3:
            celda.fill = PatternFill("solid", fgColor="FFFFFFCC")   # otra marca

    memoria = io.BytesIO()
    libro.save(memoria)
    return memoria.getvalue()


def test_analiza_tabla_y_calendario_a_la_vez():
    analisis = importer.analizar(_libro_de_prueba(), "prueba.xlsx")

    assert analisis["hoja_datos"] == "Kikis"
    assert analisis["por_tipo"] == {"Kiki": 1, "No Kiki": 1, "Marea": 3}
    assert analisis["mareas_del_calendario"] == 3
    # El amarillo no es marea, pero se reporta en lugar de desaparecer.
    assert analisis["colores_ignorados"] == {"FFFFFFCC": 1}

    kiki = next(r for r in analisis["registros"] if r["tipo"] == "Kiki")
    assert kiki["motivacion"] == "Propia" and kiki["observaciones"] == "buena"
    no_kiki = next(r for r in analisis["registros"] if r["tipo"] == "No Kiki")
    assert no_kiki["pretexto"] is None            # "N/A" no es un pretexto


def test_importar_es_idempotente(cliente):
    contenido = _libro_de_prueba()
    primero = importer.importar(contenido, "prueba.xlsx")
    assert primero["insertadas"] == 5

    segundo = importer.importar(contenido, "prueba.xlsx")
    assert segundo["insertadas"] == 0
    assert segundo["ya_existentes"] == 5


def test_simular_no_escribe(cliente):
    informe = importer.importar(_libro_de_prueba(), "prueba.xlsx", simular=True)
    assert informe["simulado"] is True
    assert informe["insertadas"] == 0
    assert cliente.get("/api/events").json()["total"] == 0


def test_modo_reemplazar_vacia_antes(cliente):
    cliente.post("/api/events", json={"fecha": "2020-01-01", "tipo": "Kiki", "calidad": 1})
    informe = importer.importar(_libro_de_prueba(), "prueba.xlsx", modo="reemplazar")
    assert informe["borradas"] == 1
    assert cliente.get("/api/events?year=2020").json()["total"] == 0


def test_csv_sin_calendario():
    csv = "fecha,pretexto,motivacion,calidad,tiempo\n2026-02-01,Rutina,Ambos,2,2\n"
    analisis = importer.analizar(csv.encode(), "datos.csv")
    assert analisis["por_tipo"]["Kiki"] == 1
    assert analisis["mareas_del_calendario"] == 0


def test_errores_claros():
    with pytest.raises(ImportError_, match="vacío"):
        importer.analizar(b"", "x.xlsx")
    with pytest.raises(ImportError_, match="Formato no soportado"):
        importer.analizar(b"algo", "x.pdf")
    with pytest.raises(ImportError_, match="cabecera"):
        importer.analizar(b"sin,nada,util\n1,2,3\n", "x.csv")


def test_endpoint_de_importacion(cliente):
    respuesta = cliente.post(
        "/api/import",
        files={"archivo": ("prueba.xlsx", _libro_de_prueba(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"modo": "combinar", "simular": "false"},
    )
    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.json()["insertadas"] == 5
    assert cliente.get("/api/import").json()["ultima_importacion"]


def test_endpoint_rechaza_formato_invalido(cliente):
    respuesta = cliente.post(
        "/api/import", files={"archivo": ("x.pdf", b"contenido", "application/pdf")}
    )
    assert respuesta.status_code == 422
